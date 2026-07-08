from __future__ import annotations

import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

from app.library_store import CompareLibraryRecord
from app.reports.xlsx_report import write_bulk_findings_xlsx


def _record(record_id: str = "abc123") -> CompareLibraryRecord:
    return CompareLibraryRecord(
        record_id=record_id,
        title="Finance Policy: Prod vs Test",
        backup_a_title="Finance Policy - Prod",
        backup_b_title="Finance Policy - Test",
        backup_a_path="C:/Backups/Finance-Prod",
        backup_b_path="C:/Backups/Finance-Test",
        saved_at="2026-07-01T09:00:00",
        record_path=f"C:/Library/Compares/{record_id}/compare.json",
        html_path=f"C:/Library/Compares/{record_id}/report.html",
        markdown_path=f"C:/Library/Compares/{record_id}/report.md",
        total_items=3,
        changed=1,
        added=1,
        removed=1,
        reviewed=0,
        actionable=3,
        ignored=0,
        source_status="Sources available",
        risk_counts={},
        diagnostics={},
    )


def _payload() -> dict:
    return {
        "inventory": [
            {
                "key": "k1",
                "name": "Enforce password history",
                "status": "Changed",
                "scope": "Computer Configuration",
                "category": "Security Options",
                "risk": "Security",
                "state_a": "24 passwords remembered",
                "state_b": "12 passwords remembered",
                "review": {
                    "status": "Under Investigation",
                    "priority": "High",
                    "owner": "J. Doe",
                    "ticket": "CHG-1042",
                    "tags": "password-policy",
                    "notes": "Confirmed intentional hardening change.",
                },
            },
            {
                "key": "k2",
                "name": "Map Network Drive",
                "status": "Added",
                "scope": "User Configuration",
                "category": "Drive Maps",
                "risk": "Preference",
                "state_a": "",
                "state_b": "Z: \\\\srv\\share",
            },
            {
                "key": "k3",
                "name": "Legacy Logon Script",
                "status": "Removed",
                "scope": "User Configuration",
                "category": "Scripts",
                "risk": "Scripts",
                "state_a": "logon.bat",
                "state_b": "",
            },
        ],
    }


class WriteBulkFindingsXlsxTests(unittest.TestCase):
    def test_writes_summary_and_findings_sheets(self) -> None:
        record = _record()
        payload = _payload()

        with TemporaryDirectory() as tmp:
            path = Path(tmp) / "findings.xlsx"
            write_bulk_findings_xlsx(path, [(record, payload)])

            workbook = load_workbook(path)
            self.assertEqual(workbook.sheetnames, ["Summary", "Findings"])

            summary = workbook["Summary"]
            self.assertEqual(summary["A1"].value, "Report")
            self.assertEqual(summary["A2"].value, record.title)
            self.assertEqual(summary["E2"].value, record.total_items)

            findings = workbook["Findings"]
            self.assertEqual(
                [cell.value for cell in findings[1]],
                [
                    "Backup A", "Backup B", "Item", "Status", "Category", "Scope",
                    "Setting A", "Setting B", "Review Status", "Priority", "Owner",
                    "Ticket/Change", "Tags", "Notes",
                ],
            )
            self.assertEqual(findings.max_row, 4)  # header + 3 items
            row2 = [cell.value for cell in findings[2]]
            self.assertEqual(row2[0], record.backup_a_title)
            self.assertEqual(row2[2], "Enforce password history")
            self.assertEqual(row2[3], "Changed")
            self.assertEqual(row2[8], "Under Investigation")
            self.assertEqual(row2[10], "J. Doe")
            self.assertEqual(row2[11], "CHG-1042")

            statuses = [findings.cell(row=r, column=4).value for r in range(2, 5)]
            self.assertEqual(statuses, ["Changed", "Missing in A", "Missing in B"])

            self.assertEqual(findings.auto_filter.ref, "A1:N4")
            self.assertEqual(findings.freeze_panes, "A2")

    def test_multiple_reports_combine_into_one_findings_sheet(self) -> None:
        record_a = _record("a1")
        record_b = _record("b2")
        payload = _payload()

        with TemporaryDirectory() as tmp:
            path = Path(tmp) / "findings.xlsx"
            write_bulk_findings_xlsx(path, [(record_a, payload), (record_b, payload)])

            workbook = load_workbook(path)
            findings = workbook["Findings"]
            self.assertEqual(findings.max_row, 7)  # header + 3 + 3

            summary = workbook["Summary"]
            self.assertEqual(summary.max_row, 3)  # header + 2 records


if __name__ == "__main__":
    unittest.main()
