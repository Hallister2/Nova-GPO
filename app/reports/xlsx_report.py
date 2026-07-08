from __future__ import annotations

from pathlib import Path
from typing import TYPE_CHECKING, Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.reports.compare_report import _status_label

if TYPE_CHECKING:
    from app.library_store import CompareLibraryRecord

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(color="FFFFFF", bold=True)
STATUS_FILLS = {
    "Missing in A": PatternFill("solid", fgColor="DCFCE7"),
    "Missing in B": PatternFill("solid", fgColor="FEE2E2"),
    "Changed": PatternFill("solid", fgColor="FEF3C7"),
    "Different": PatternFill("solid", fgColor="FEF3C7"),
}

FINDINGS_HEADERS = [
    "Backup A",
    "Backup B",
    "Item",
    "Status",
    "Category",
    "Scope",
    "Setting A",
    "Setting B",
    "Review Status",
    "Priority",
    "Owner",
    "Ticket/Change",
    "Tags",
    "Notes",
]

_STATUS_COLUMN = 4

SUMMARY_HEADERS = [
    "Report",
    "Backup A",
    "Backup B",
    "Saved At",
    "Total Items",
    "Actionable",
    "Added",
    "Changed",
    "Removed",
    "Reviewed",
    "Ignored",
]


def write_bulk_findings_xlsx(
    path: str | Path,
    entries: list[tuple["CompareLibraryRecord", dict[str, Any]]],
) -> None:
    """Write a single workbook covering every selected saved compare report.

    entries: (record, payload) pairs where payload is the dict returned by
    ``load_compare_record_payload`` for that record.
    """
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    _write_summary_sheet(summary_sheet, entries)

    findings_sheet = workbook.create_sheet("Findings")
    _write_findings_sheet(findings_sheet, entries)

    workbook.save(str(path))


def _write_summary_sheet(
    sheet: Worksheet,
    entries: list[tuple["CompareLibraryRecord", dict[str, Any]]],
) -> None:
    _write_header_row(sheet, SUMMARY_HEADERS)
    for record, _payload in entries:
        sheet.append([
            record.title,
            record.backup_a_title,
            record.backup_b_title,
            record.saved_at,
            record.total_items,
            record.actionable,
            record.added,
            record.changed,
            record.removed,
            record.reviewed,
            record.ignored,
        ])
    _finalize_sheet(sheet, len(SUMMARY_HEADERS), len(entries))


def _write_findings_sheet(
    sheet: Worksheet,
    entries: list[tuple["CompareLibraryRecord", dict[str, Any]]],
) -> None:
    _write_header_row(sheet, FINDINGS_HEADERS)

    row_count = 0
    for record, payload in entries:
        items = payload.get("inventory") or payload.get("items") or []
        for item in items:
            status = _status_label(str(item.get("status", "")))
            review = item.get("review") or {}
            sheet.append([
                record.backup_a_title,
                record.backup_b_title,
                item.get("name", ""),
                status,
                item.get("category", ""),
                item.get("scope", ""),
                item.get("state_a", ""),
                item.get("state_b", ""),
                review.get("status", ""),
                review.get("priority", ""),
                review.get("owner", ""),
                review.get("ticket", ""),
                review.get("tags", ""),
                review.get("notes", ""),
            ])
            row_count += 1
            fill = STATUS_FILLS.get(status)
            if fill is not None:
                sheet.cell(row=sheet.max_row, column=_STATUS_COLUMN).fill = fill

    _finalize_sheet(sheet, len(FINDINGS_HEADERS), row_count)


def _write_header_row(sheet: Worksheet, headers: list[str]) -> None:
    sheet.append(headers)
    for column_index in range(1, len(headers) + 1):
        cell = sheet.cell(row=1, column=column_index)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A2"


def _finalize_sheet(sheet: Worksheet, column_count: int, row_count: int) -> None:
    last_column_letter = get_column_letter(column_count)
    sheet.auto_filter.ref = f"A1:{last_column_letter}{max(row_count + 1, 1)}"
    for column_index in range(1, column_count + 1):
        letter = get_column_letter(column_index)
        longest = max(
            (len(str(cell.value)) for cell in sheet[letter] if cell.value is not None),
            default=10,
        )
        sheet.column_dimensions[letter].width = min(max(longest + 2, 12), 60)
